"""output/drafts/ 内の下書きファイルのフロントマターを集計し、output/posting_log.xlsx を更新する。

下書きファイルのフォーマットは output/drafts/README.md を参照。
このスクリプトは集計のみを行う（下書きの文章そのものはClaudeが対話の中で作成する）。
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DRAFTS_DIR = ROOT / "output" / "drafts"
LOG_PATH = ROOT / "output" / "posting_log.xlsx"

FRONT_MATTER_RE = re.compile(r"^---\n(.*?)\n---\n", re.DOTALL)
BOOL_TRUE = {"true", "yes"}

COLUMNS = [
    ("ファイル名", "file_name"),
    ("商品名", "product_name"),
    ("商品URL", "product_url"),
    ("作成日", "created_date"),
    ("楽天ROOM投稿済み", "rakuten_room_posted"),
    ("Instagram投稿済み", "instagram_posted"),
    ("X投稿済み", "x_posted"),
]


def parse_front_matter(text: str) -> dict:
    match = FRONT_MATTER_RE.match(text)
    if not match:
        return {}
    fields = {}
    for line in match.group(1).splitlines():
        if ":" not in line:
            continue
        key, _, value = line.partition(":")
        fields[key.strip()] = value.strip()
    return fields


def to_bool_label(value: str) -> str:
    return "済" if value.lower() in BOOL_TRUE else "未"


def collect_rows() -> list[dict]:
    rows = []
    for path in sorted(DRAFTS_DIR.glob("*.md")):
        if path.name == "README.md" or path.name.startswith("_"):
            continue
        fields = parse_front_matter(path.read_text(encoding="utf-8"))
        if not fields:
            print(f"[警告] フロントマターが見つかりません: {path.name}", file=sys.stderr)
            continue
        rows.append(
            {
                "file_name": path.name,
                "product_name": fields.get("product_name", ""),
                "product_url": fields.get("product_url", ""),
                "created_date": fields.get("created_date", ""),
                "rakuten_room_posted": to_bool_label(fields.get("rakuten_room_posted", "false")),
                "instagram_posted": to_bool_label(fields.get("instagram_posted", "false")),
                "x_posted": to_bool_label(fields.get("x_posted", "false")),
            }
        )
    return rows


def write_log(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "投稿管理ログ"

    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (_, key) in enumerate(COLUMNS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row[key])

    for col_idx, (header, key) in enumerate(COLUMNS, start=1):
        widths = [len(header)] + [len(str(row[key])) for row in rows]
        ws.column_dimensions[get_column_letter(col_idx)].width = max(widths) + 4

    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(LOG_PATH)


def main() -> None:
    if not DRAFTS_DIR.exists():
        print(f"[エラー] {DRAFTS_DIR} が見つかりません", file=sys.stderr)
        sys.exit(1)

    rows = collect_rows()
    write_log(rows)
    print(f"{len(rows)}件の下書きを集計し、{LOG_PATH.relative_to(ROOT)} を更新しました。")


if __name__ == "__main__":
    main()
